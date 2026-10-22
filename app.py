from __future__ import annotations
import io, os, re, sqlite3, zipfile
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from copy import copy

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from docx import Document

ROOT = Path(__file__).parent
ASSETS = ROOT / "assets" / "templates"
DB_PATH = Path(os.environ.get("TRANSPORT_DB_PATH", str(ROOT / "data" / "transport.db")))
DB_PATH.parent.mkdir(parents=True, exist_ok=True)

st.set_page_config(page_title="YANAH LOGISTICS MANAGEMENT SYSTEM", page_icon="🚚", layout="wide", initial_sidebar_state="expanded")
st.markdown("""
<style>
:root { --ink:#172554; --blue:#2563eb; --violet:#7c3aed; --teal:#0891b2; --mint:#d1fae5; }
.stApp { background: linear-gradient(135deg, #0b3d91 0%, #1456b8 48%, #087eaa 100%); color: #ffffff; }
.stApp p, .stApp h1, .stApp h2, .stApp h3, .stApp h4, .stApp label, .stApp [data-testid="stMarkdownContainer"] { color: #ffffff; }
[data-testid="stDataFrame"], [data-testid="stTable"] { background:#ffffff; color:#172554; border-radius:12px; }
[data-testid="stSidebar"] { background:linear-gradient(180deg,#f1f5ff 0%,#f5f3ff 55%,#ecfeff 100%); border-right:1px solid #dbeafe; }
[data-testid="stSidebar"] * { color:#172554 !important; }
[data-testid="stSidebar"] [data-testid="stRadio"] > label { font-weight:800; font-size:.82rem; color:#475569 !important; margin-bottom:12px; }
[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] { gap:10px; }
[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] > label { background:rgba(255,255,255,.9); border:1px solid #dbeafe; border-radius:13px; padding:12px 14px; margin:0; box-shadow:0 3px 10px rgba(30,64,175,.05); min-height:46px; }
[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] > label:hover { background:#e0e7ff; border-color:#818cf8; }
[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] > label:has(input:checked) { background:linear-gradient(100deg,#dbeafe,#ede9fe); border:1px solid #818cf8; box-shadow:inset 4px 0 0 #4f46e5,0 4px 12px rgba(79,70,229,.10); }
[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] > label > div:first-child { display:none; }
.block-container { padding-top: 1.2rem; padding-bottom: 2rem; max-width: 1400px; }
.hero { background: linear-gradient(110deg,#1d4ed8,#6d28d9 58%,#0891b2); color:white; border-radius:22px; padding:24px 28px; box-shadow:0 14px 35px rgba(37,99,235,.18); margin-bottom:18px; }
.hero h1 { color:white; margin:0; font-size:2rem; }
.hero p { color:#e0f2fe; margin:.35rem 0 0 0; }
.hero.booking-hero { background:linear-gradient(115deg,#eff6ff 0%,#f5f3ff 52%,#ecfeff 100%); color:#172554; border:1px solid #c7d2fe; box-shadow:0 10px 28px rgba(79,70,229,.09); }
.hero.booking-hero h1 { color:#1e3a8a; }
.hero.booking-hero p { color:#475569; }
body:has(.booking-form-marker) [data-testid="stForm"] { background:#ffffff; border:1px solid #c7d2fe; box-shadow:0 6px 18px rgba(30,64,175,.07); }
/* Dashboard booking form: vivid but readable, with a soft glass-card finish */
body:has(.dashboard-form-marker) [data-testid="stForm"] {
  background:linear-gradient(135deg,#ffffff 0%,#eff6ff 48%,#f5f3ff 100%) !important;
  border:1px solid #93c5fd !important; border-radius:22px !important;
  padding:22px !important; box-shadow:0 14px 34px rgba(8,35,90,.24), inset 0 1px 0 #fff !important;
}
body:has(.dashboard-form-marker) [data-testid="stForm"] label,
body:has(.dashboard-form-marker) [data-testid="stForm"] [data-testid="stMarkdownContainer"],
body:has(.dashboard-form-marker) [data-testid="stForm"] p { color:#172554 !important; font-weight:650 !important; }
body:has(.dashboard-form-marker) [data-testid="stForm"] input,
body:has(.dashboard-form-marker) [data-testid="stForm"] textarea,
body:has(.dashboard-form-marker) [data-testid="stForm"] [data-baseweb="select"] { background:#fff !important; color:#172554 !important; border-color:#93c5fd !important; border-radius:10px !important; }
body:has(.dashboard-form-marker) [data-testid="stForm"] [data-testid="stFormSubmitButton"] button { background:linear-gradient(100deg,#2563eb,#7c3aed) !important; color:#fff !important; border:0 !important; border-radius:13px !important; box-shadow:0 6px 14px rgba(79,70,229,.25); }
.dashboard-form-heading { background:linear-gradient(100deg,#dbeafe,#ede9fe 58%,#cffafe); border:1px solid #bfdbfe; border-radius:16px; padding:13px 18px; margin:8px 0 10px; box-shadow:0 6px 16px rgba(30,64,175,.12); }
.dashboard-form-heading h3 { color:#172554 !important; margin:0; }
.dashboard-form-heading p { color:#475569 !important; margin:4px 0 0; }
body:has(.booking-form-marker) [data-testid="stTextInput"] input,
body:has(.booking-form-marker) [data-testid="stNumberInput"] input,
body:has(.booking-form-marker) [data-testid="stSelectbox"] div[data-baseweb="select"] { background:#fff !important; color:#172554 !important; border-color:#cbd5e1 !important; }
body:has(.booking-form-marker) [data-testid="stForm"] label { color:#1e293b !important; font-weight:650; }
.metric-card { background:rgba(255,255,255,.92); border:1px solid #dbeafe; border-radius:18px; padding:16px 18px; box-shadow:0 8px 22px rgba(30,64,175,.06); min-height:108px; }
.metric-label { color:#64748b; font-size:.82rem; font-weight:700; text-transform:uppercase; letter-spacing:.04em; }
.metric-value { color:#1e3a8a; font-size:1.8rem; font-weight:800; margin-top:5px; }
.section-card { background:rgba(255,255,255,.96); color:#172554; border:1px solid #bfdbfe; border-radius:18px; padding:18px; box-shadow:0 10px 25px rgba(3,20,60,.18); }
.section-card h3, .section-card p { color:#172554 !important; }
.stButton > button, .stDownloadButton > button { border-radius:12px; font-weight:700; min-height:44px; }
.stButton > button[kind="primary"] { background:linear-gradient(100deg,#2563eb,#7c3aed); border:0; }
.stDownloadButton > button { background:linear-gradient(100deg,#0e7490,#2563eb); color:white; border:0; }
[data-testid="stFileUploader"] { background:#fff; border:2px dashed #a5b4fc; border-radius:16px; padding:10px; }
div[data-testid="stForm"] { background:rgba(255,255,255,.96); color:#172554; border:1px solid #bfdbfe; padding:16px; border-radius:16px; box-shadow:0 8px 22px rgba(3,20,60,.12); }
div[data-testid="stForm"] label { color:#172554 !important; font-weight:650; }
div[data-testid="stForm"] input, div[data-testid="stForm"] textarea, div[data-testid="stForm"] [data-baseweb="select"] { background:#fff !important; color:#172554 !important; }
.small-muted { color:#64748b; font-size:.88rem; }
</style>
""", unsafe_allow_html=True)

PRODUCTS = {
    "CHROME CONCENTRATE": "chrome_concentrate.docx",
    "HIGH CARBON FERROCHROME": "high_carbon_ferrochrome.docx",
    "LITHIUM": "lithium.docx",
    "SILICA SANDS & QUARTZ SANDS": "silica_quartz_sands.docx",
}
TRUCK_HEADERS = ["Booking reff", "SN", "Truck Reg.", "Trailor Reg.", "Container No.", "Seal", "Driver Name", "Phone No.", "Licence No.", "Passport No.", "Capacity", "Tare weight (KG)"]
BOOKING_HEADERS = ["Business Ref.", "Booking reff", "Vessel", "ETB", "Transporter", "SN", "Truck Reg.", "Trailor Reg.", "Container No.", "Seal", "Driver Name", "Phone No.", "Licence No.", "Passport No.", "Capacity", "Loading Site", "Booking Date"]
ALIASES = {
    "booking reff":"Booking reff", "booking ref":"Booking reff", "booking number":"Booking reff",
    "sn":"SN", "serial number":"SN", "serial no":"SN",
    "truck reg.":"Truck Reg.", "truck reg":"Truck Reg.", "horse reg":"Truck Reg.", "horse registration":"Truck Reg.",
    "trailor reg.":"Trailor Reg.", "trailor reg":"Trailor Reg.", "trailer reg.":"Trailor Reg.", "trailer reg":"Trailor Reg.", "trailer registration":"Trailor Reg.",
    "container no.":"Container No.", "container no":"Container No.", "container number":"Container No.",
    "seal":"Seal", "seal no":"Seal", "seal number":"Seal", "driver name":"Driver Name", "phone no.":"Phone No.", "phone no":"Phone No.", "phone":"Phone No.", "contact":"Phone No.",
    "licence no.":"Licence No.", "licence no":"Licence No.", "license no.":"Licence No.", "license no":"Licence No.",
    "passport no.":"Passport No.", "passport no":"Passport No.", "capacity":"Capacity", "tare weight (kg)":"Tare weight (KG)", "tare weight":"Tare weight (KG)",
    "transporter":"Transporter", "fleet":"Transporter", "fleet name":"Transporter"
}

def now_local():
    return datetime.now(ZoneInfo("Africa/Harare"))

def clean(v):
    if v is None or (isinstance(v, float) and pd.isna(v)): return ""
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def norm_reg(v): return re.sub(r"\s+", "", clean(v)).upper()
def norm_header(v): return re.sub(r"\s+", " ", clean(v).lower()).strip().rstrip(":")
def canon(v):
    n = norm_header(v)
    return ALIASES.get(n, clean(v))
def map_transporter(v):
    s = clean(v)
    u = s.upper()
    if "OP FLEET A" in u or u == "OP (A)": return "OP (A)"
    if "OP FLEET B" in u or u == "OP (B)": return "OP (B)"
    return s

def db():
    con = sqlite3.connect(DB_PATH, check_same_thread=False)
    con.row_factory = sqlite3.Row
    con.execute("""CREATE TABLE IF NOT EXISTS bookings (
      booking_no TEXT PRIMARY KEY, business_ref TEXT, vessel TEXT, etb TEXT, loading_point TEXT,
      expected_trucks INTEGER, booking_date TEXT, product TEXT, description TEXT, consignor TEXT,
      consignee TEXT, weight_mass TEXT, created_at TEXT)""")
    con.execute("""CREATE TABLE IF NOT EXISTS trucks (
      id INTEGER PRIMARY KEY AUTOINCREMENT, booking_no TEXT, sn INTEGER, truck_reg TEXT, trailer_reg TEXT,
      container_no TEXT, seal TEXT, driver_name TEXT, phone_no TEXT, licence_no TEXT, passport_no TEXT,
      capacity TEXT, tare_weight TEXT, transporter TEXT, batch_uploaded_at TEXT, manifest_generated INTEGER DEFAULT 0,
      UNIQUE(booking_no, truck_reg))""")
    con.commit()
    return con

con = db()

def auth_gate():
    """Fail-closed password gate. APP_PASSWORD must be configured in Streamlit Secrets."""
    import hmac
    try:
        expected = str(st.secrets.get("APP_PASSWORD", "")).strip()
    except Exception:
        expected = ""

    if not expected:
        st.markdown(
            '<div class="hero"><h1>🔐 YANAH LOGISTICS MANAGEMENT SYSTEM</h1>'
            '<p>Private access is not configured yet.</p></div>',
            unsafe_allow_html=True,
        )
        st.error("For security, the system is locked until the administrator configures APP_PASSWORD in Streamlit Secrets.")
        st.info("Open your app in Streamlit Community Cloud → Settings / Manage app → Secrets, add APP_PASSWORD, then reboot the app.")
        return False

    if st.session_state.get("authenticated"):
        return True

    st.markdown(
        '<div class="hero"><h1>🔐 YANAH LOGISTICS MANAGEMENT SYSTEM</h1>'
        '<p>Secure administrator sign-in</p></div>',
        unsafe_allow_html=True,
    )
    st.markdown(
        '<div class="section-card"><h3>Welcome back</h3>'
        '<p>Sign in to access private bookings, truck records and road manifests.</p></div>',
        unsafe_allow_html=True,
    )
    with st.form("login_form", clear_on_submit=True):
        password = st.text_input("Administrator password", type="password", autocomplete="current-password")
        submitted = st.form_submit_button("Sign in securely", type="primary", use_container_width=True)
    if submitted:
        if hmac.compare_digest(password, expected):
            st.session_state.authenticated = True
            st.rerun()
        st.error("Incorrect password. Please try again.")
    return False

if not auth_gate():
    st.stop()

# Navigation
st.sidebar.markdown("# 🚚 YANAH LOGISTICS")
st.sidebar.caption("BOOKINGS • TRUCKS • ROAD MANIFESTS")
if st.sidebar.button("🔒 Log out", use_container_width=True):
    st.session_state.pop("authenticated", None)
    st.rerun()
nav_options = ["Dashboard", "Create a booking", "Add trucks & generate", "Booking records", "Templates & setup"]
nav_target = st.session_state.pop("nav_target", "Dashboard")
page = st.sidebar.radio("Navigate", nav_options, index=nav_options.index(nav_target) if nav_target in nav_options else 0, label_visibility="collapsed")
st.sidebar.markdown("---")
st.sidebar.caption("Bright, mobile-friendly workspace")


def hero(title, subtitle):
    extra_class = " booking-hero" if title == "📝 Create a booking" else ""
    st.markdown(f'<div class="hero{extra_class}"><h1>{title}</h1><p>{subtitle}</p></div>', unsafe_allow_html=True)

def metric(label, value, note=""):
    st.markdown(f'<div class="metric-card"><div class="metric-label">{label}</div><div class="metric-value">{value}</div><div class="small-muted">{note}</div></div>', unsafe_allow_html=True)

def get_bookings(): return con.execute("SELECT * FROM bookings ORDER BY created_at DESC").fetchall()
def get_booking(no): return con.execute("SELECT * FROM bookings WHERE booking_no=?", (no,)).fetchone()
def get_trucks(no): return con.execute("SELECT * FROM trucks WHERE booking_no=? ORDER BY sn, id", (no,)).fetchall()

def parse_truck_file(file_bytes, filename):
    ext = Path(filename).suffix.lower()
    if filename.lower().endswith(".xlt.xls") or ext == ".xls":
        raw = pd.read_excel(io.BytesIO(file_bytes), header=None, dtype=object, engine="xlrd").fillna("")
    elif ext in (".xlsx", ".xlsm"):
        raw = pd.read_excel(io.BytesIO(file_bytes), header=None, dtype=object, engine="openpyxl").fillna("")
    elif ext == ".csv":
        raw = pd.read_csv(io.BytesIO(file_bytes), header=None, dtype=object).fillna("")
    else: raise ValueError("Use .xls, .xlsx, .xlsm or .csv truck lists.")
    # Locate a row with both truck and trailer headings.
    header_idx = None
    for i in range(min(len(raw), 80)):
        vals = [norm_header(x) for x in raw.iloc[i].tolist()]
        has_truck = any(v in ("truck reg", "truck reg.", "horse reg", "horse registration") for v in vals)
        has_trailer = any(v in ("trailor reg", "trailor reg.", "trailer reg", "trailer reg.", "trailer registration") for v in vals)
        if has_truck and has_trailer:
            header_idx = i; break
    if header_idx is None: raise ValueError("Could not find Truck Reg. and Trailer/Trailor Reg. column headings.")
    headers = [canon(x) for x in raw.iloc[header_idx].tolist()]
    df = raw.iloc[header_idx+1:].copy()
    df.columns = headers
    # Fleet labels in the rows above and inside data.
    current_fleet = ""
    for i in range(header_idx):
        joined = " ".join(clean(x) for x in raw.iloc[i].tolist()).upper()
        if "OP FLEET A" in joined: current_fleet = "OP (A)"
        elif "OP FLEET B" in joined: current_fleet = "OP (B)"
    fleet_values=[]
    for _, row in df.iterrows():
        joined = " ".join(clean(x) for x in row.tolist()).upper()
        if "OP FLEET A" in joined: current_fleet = "OP (A)"
        elif "OP FLEET B" in joined: current_fleet = "OP (B)"
        explicit = next((clean(row.get(c)) for c in df.columns if canon(c)=="Transporter" and clean(row.get(c))), "")
        fleet_values.append(map_transporter(explicit or current_fleet))
    df["_Transporter"] = fleet_values
    truck_col = next((c for c in df.columns if canon(c)=="Truck Reg."), None)
    if not truck_col: raise ValueError("Truck registration column was not found.")
    df = df[df[truck_col].map(lambda x: bool(norm_reg(x)))].copy()
    # Normalize duplicate column labels by keeping first occurrence.
    df = df.loc[:, ~pd.Index(df.columns).duplicated()]
    return df

def source_value(row, *fields):
    for f in fields:
        col = next((c for c in row.index if canon(c)==f), None)
        if col is not None and clean(row.get(col)): return clean(row.get(col))
    return ""

def validate_container(value):
    v = re.sub(r"\s+", "", clean(value)).upper()
    if not v: return "missing"
    return "valid" if re.fullmatch(r"[A-Z]{4}\d{7}", v) else "invalid"

def truck_record(row, booking_no, sn, transporter_default=""):
    reg = norm_reg(source_value(row,"Truck Reg."))
    trailer = norm_reg(source_value(row,"Trailor Reg."))
    container = re.sub(r"\s+", "", source_value(row,"Container No.")).upper()
    seal = re.sub(r"\s+", "", source_value(row,"Seal")).upper()
    transporter = map_transporter(source_value(row,"Transporter") or source_value(row,"_Transporter") or transporter_default)
    return {"booking_no":booking_no,"sn":sn,"truck_reg":reg,"trailer_reg":trailer,"container_no":container,"seal":seal,
      "driver_name":source_value(row,"Driver Name"),"phone_no":source_value(row,"Phone No."),"licence_no":source_value(row,"Licence No."),
      "passport_no":source_value(row,"Passport No."),"capacity":source_value(row,"Capacity"),"tare_weight":source_value(row,"Tare weight (KG)"),"transporter":transporter}

def copy_row_style(ws, src_row, dst_row, max_col):
    if src_row == dst_row: return
    for c in range(1, max_col+1):
        a, b = ws.cell(src_row,c), ws.cell(dst_row,c)
        if a.has_style: b._style = copy(a._style)
        if a.number_format: b.number_format = a.number_format
        if a.alignment: b.alignment = copy(a.alignment)
        if a.protection: b.protection = copy(a.protection)
    if ws.row_dimensions[src_row].height: ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height

def workbook_bytes(template_path, rows, headers, row_builder, clear_from=2):
    wb = load_workbook(template_path)
    ws = wb[wb.sheetnames[0]]
    for r in range(clear_from, ws.max_row+1):
        for c in range(1, ws.max_column+1): ws.cell(r,c).value = None
    max_col = len(headers)
    for i, record in enumerate(rows, start=0):
        r = clear_from+i
        if r > ws.max_row: copy_row_style(ws, clear_from, r, max_col)
        vals = row_builder(record)
        for c, h in enumerate(headers, start=1): ws.cell(r,c).value = vals.get(h, "")
    out=io.BytesIO(); wb.save(out); return out.getvalue()

def make_truck_list(booking, truck_rows):
    template = ASSETS / "truck_list_template.xlsx"
    # Use exact supplied 12-column workbook and its row formatting.
    def build(t):
        return {"Booking reff":booking["booking_no"],"SN":t["sn"],"Truck Reg.":t["truck_reg"],"Trailor Reg.":t["trailer_reg"],
          "Container No.":t["container_no"],"Seal":t["seal"],"Driver Name":t["driver_name"],"Phone No.":t["phone_no"],
          "Licence No.":t["licence_no"],"Passport No.":t["passport_no"],"Capacity":t["capacity"],"Tare weight (KG)":t["tare_weight"]}
    # Preserve the exact header row; use row 2's style for additional rows.
    return workbook_bytes(template, truck_rows, TRUCK_HEADERS, build)

def make_booking_list(booking, all_trucks):
    template = ASSETS / "booking_template.xlsx"
    def build(t):
        return {"Business Ref.":booking["business_ref"],"Booking reff":booking["booking_no"],"Vessel":booking["vessel"],"ETB":booking["etb"],
          "Transporter":t["transporter"],"SN":t["sn"],"Truck Reg.":t["truck_reg"],"Trailor Reg.":t["trailer_reg"],"Container No.":t["container_no"],
          "Seal":t["seal"],"Driver Name":t["driver_name"],"Phone No.":t["phone_no"],"Licence No.":t["licence_no"],"Passport No.":t["passport_no"],
          "Capacity":t["capacity"],"Loading Site":booking["loading_point"],"Booking Date":booking["booking_date"]}
    return workbook_bytes(template, all_trucks, BOOKING_HEADERS, build)

def replace_cell(cell, value):
    value = clean(value)
    if not cell.paragraphs:
        cell.text = value; return
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = value
        for run in p.runs[1:]: run.text = ""
    else: p.add_run(value)
    for extra in cell.paragraphs[1:]:
        for run in extra.runs: run.text = ""

def replace_text_in_doc(doc, old, new):
    for p in doc.paragraphs:
        if old in p.text:
            for run in p.runs:
                if old in run.text: run.text = run.text.replace(old,new)
    for table in doc.tables:
        for row in table.rows:
            # Avoid repeating merged cells by their XML identity.
            seen=set()
            for cell in row.cells:
                key=id(cell._tc)
                if key in seen: continue
                seen.add(key)
                for p in cell.paragraphs:
                    for run in p.runs:
                        if old in run.text: run.text=run.text.replace(old,new)

def make_manifest(booking, truck, product_template_path):
    doc = Document(product_template_path)
    if not doc.tables or len(doc.tables[0].rows) < 13: raise ValueError("Road Manifest template structure does not match the supplied template.")
    table = doc.tables[0]
    # Based on the supplied ROAD MANIFEST document; inspect every draft before official use.
    replace_cell(table.rows[6].cells[3], truck["truck_reg"])
    replace_cell(table.rows[6].cells[7], truck["trailer_reg"])
    replace_cell(table.rows[7].cells[3], truck["container_no"])
    replace_cell(table.rows[8].cells[3], truck["seal"])
    # Shared cargo fields from booking profile.
    replace_cell(table.rows[11].cells[4], booking["weight_mass"])
    replace_cell(table.rows[11].cells[5], booking["description"] or booking["product"])
    replace_cell(table.rows[11].cells[9], booking["consignor"])
    replace_cell(table.rows[11].cells[10], booking["consignee"])
    # Put the driver name on the same horizontal line as the "Name of Driver/Transporter" label,
    # between that label and "Signature". The old placement used a blank paragraph below the label.
    driver_cell = table.rows[12].cells[0]
    driver_paragraph = next((p for p in driver_cell.paragraphs if "Name of Driver/Transporter" in p.text and "Signature" in p.text), None)
    if driver_paragraph is not None:
        # The supplied template splits the label and tabs across several runs. Keep the original
        # label formatting and signature label, and insert the name at the existing tab position.
        signature_run = next((i for i, run in enumerate(driver_paragraph.runs) if run.text.strip() == "Signature"), None)
        tab_run = next((i for i, run in enumerate(driver_paragraph.runs) if "\t" in run.text), None)
        if tab_run is not None:
            driver_paragraph.runs[tab_run].text = "\t" + (truck["driver_name"] or "") + "\t"
        elif signature_run is not None:
            driver_paragraph.runs[signature_run].text = "\t" + (truck["driver_name"] or "") + "\tSignature"
        else:
            driver_paragraph.add_run("\t" + (truck["driver_name"] or ""))
    else:
        # Safe fallback: use the existing paragraph immediately after the label rather than
        # adding another paragraph that could shift the form layout.
        driver_cell.paragraphs[7].add_run(truck["driver_name"] or "")
    # Replace only the date text in the certification paragraph, preserving the Date and Place labels.
    try:
        date_text = datetime.strptime(booking["booking_date"], "%d.%m.%y").strftime("%d %B, %Y").upper()
    except Exception:
        date_text = booking["booking_date"]
    for table_obj in doc.tables:
        seen_cells = set()
        for row_obj in table_obj.rows:
            for cell_obj in row_obj.cells:
                if id(cell_obj._tc) in seen_cells: continue
                seen_cells.add(id(cell_obj._tc))
                for para in cell_obj.paragraphs:
                    if para.text.strip().startswith("Date ") and "Place" in para.text:
                        # In the supplied template, the date is split over runs (12 / OCTOBER / 202 / 6).
                        date_runs = [i for i, run in enumerate(para.runs) if run.text.strip() in ("12", "OCTOBER,", "202", "6") or "OCTOBER" in run.text]
                        if date_runs:
                            first = date_runs[0]
                            para.runs[first].text = date_text
                            for ri in date_runs[1:]: para.runs[ri].text = ""
    out=io.BytesIO(); doc.save(out); return out.getvalue()

def make_manifest_zip(booking, trucks):
    filename=PRODUCTS.get(booking["product"], PRODUCTS["HIGH CARBON FERROCHROME"])
    template_path=ASSETS / filename
    out=io.BytesIO(); issues=[]; count=0
    with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as zf:
        for t in trucks:
            reg=t["truck_reg"]
            if not reg: continue
            status=validate_container(t["container_no"])
            if status=="invalid":
                issues.append(f"SN {t['sn']} / {reg}: invalid container number {t['container_no']!r}; review before official use.")
            if status=="missing": issues.append(f"SN {t['sn']} / {reg}: container number is missing.")
            data=make_manifest(booking,t,template_path)
            safe=re.sub(r"[^A-Za-z0-9_-]+","_",reg)
            zf.writestr(f"{safe}.docx",data); count+=1
    return out.getvalue(), count, issues

if page == "Dashboard":
    hero("YANAH LOGISTICS MANAGEMENT SYSTEM", "Manage bookings, truck batches, and road manifests in one place.")
    bookings = get_bookings()
    total_trucks = con.execute("SELECT COUNT(*) FROM trucks").fetchone()[0]
    generated = con.execute("SELECT COUNT(*) FROM trucks WHERE manifest_generated=1").fetchone()[0]
    open_bookings = sum(1 for b in bookings if con.execute("SELECT COUNT(*) FROM trucks WHERE booking_no=?",(b['booking_no'],)).fetchone()[0] < b['expected_trucks'])
    cols=st.columns(4)
    with cols[0]: metric("Bookings",len(bookings),"Saved booking records")
    with cols[1]: metric("Trucks received",total_trucks,"Unique truck records")
    with cols[2]: metric("Manifests generated",generated,"Marked complete")
    with cols[3]: metric("Bookings in progress",open_bookings,"Below expected truck count")
    st.markdown('<div class="dashboard-form-marker"></div><div class="dashboard-form-heading"><h3>✦ Create a booking</h3><p>Enter the details below, then press Save Booking at the bottom. The booking is saved directly from this Dashboard.</p></div>', unsafe_allow_html=True)
    with st.form("dashboard_create_booking_form"):
        d1,d2,d3=st.columns(3)
        with d1: dash_booking_no=st.text_input("Booking Number *",placeholder="e.g. BEW600035000",key="dash_booking_no")
        with d2: dash_business_ref=st.text_input("Business Reference *",placeholder="e.g. OPSE26100002",key="dash_business_ref")
        with d3: dash_expected=st.number_input("Expected number of trucks *",min_value=1,max_value=1000,value=20,step=1,key="dash_expected")
        d4,d5,d6=st.columns(3)
        with d4: dash_vessel=st.text_input("Vessel",key="dash_vessel")
        with d5: dash_etb=st.text_input("ETB",key="dash_etb")
        with d6: dash_loading=st.text_input("Loading Point / Mine *",placeholder="e.g. MUTOKO",key="dash_loading")
        dash_product=st.selectbox("ROAD MANIFEST category *",list(PRODUCTS.keys()),index=1,key="dash_product")
        d7,d8=st.columns(2)
        with d7: dash_description=st.text_input("Description of goods",value="HIGH CARBON FERROCHROME",key="dash_description")
        with d8: dash_weight=st.text_input("Weight / Mass",placeholder="e.g. 28 MT",key="dash_weight")
        d9,d10=st.columns(2)
        with d9: dash_consignor=st.text_input("Consignor",key="dash_consignor")
        with d10: dash_consignee=st.text_input("Consignee",key="dash_consignee")
        dash_save=st.form_submit_button("✦ Save Booking",type="primary",use_container_width=True)
    if dash_save:
        dash_no=clean(dash_booking_no).upper()
        if not dash_no or not clean(dash_business_ref) or not clean(dash_loading):
            st.error("Please complete Booking Number, Business Reference and Loading Point.")
        elif get_booking(dash_no):
            st.error("That booking number already exists. Open it under Booking records instead of creating it again.")
        else:
            con.execute("INSERT INTO bookings VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",(dash_no,clean(dash_business_ref).upper(),clean(dash_vessel),clean(dash_etb),clean(dash_loading).upper(),int(dash_expected),"",dash_product,clean(dash_description) or dash_product,clean(dash_consignor),clean(dash_consignee),clean(dash_weight),now_local().isoformat(timespec="seconds")))
            con.commit()
            st.success(f"Booking {dash_no} created successfully from the Dashboard. You can now upload truck lists.")
            st.rerun()

    st.markdown("### Your workspace")
    a,b,c=st.columns(3)
    with a:
        st.markdown('<div class="section-card"><h3>🧾 Booking lists</h3><p>Create a booking once, keep its date fixed, and build the Excel booking list as trucks arrive.</p></div>',unsafe_allow_html=True)
        if st.button("＋ Create a booking", use_container_width=True): st.session_state.nav_target="Create a booking"; st.rerun()
    with b:
        st.markdown('<div class="section-card"><h3>🚛 Truck batches</h3><p>Upload each batch, validate container numbers, and export the matching truck-list workbook.</p></div>',unsafe_allow_html=True)
        if st.button("Upload trucks", use_container_width=True): st.session_state.nav_target="Add trucks & generate"; st.rerun()
    with c:
        st.markdown('<div class="section-card"><h3>📄 Road manifests</h3><p>Choose one of four cargo categories and generate one Word manifest per truck.</p></div>',unsafe_allow_html=True)
        if st.button("View bookings", use_container_width=True): st.session_state.nav_target="Booking records"; st.rerun()
    if bookings:
        st.markdown("### Recent bookings")
        rows=[]
        for b in bookings[:10]:
            n=con.execute("SELECT COUNT(*) FROM trucks WHERE booking_no=?",(b['booking_no'],)).fetchone()[0]
            m=con.execute("SELECT COUNT(*) FROM trucks WHERE booking_no=? AND manifest_generated=1",(b['booking_no'],)).fetchone()[0]
            rows.append({"Booking Number":b['booking_no'],"Business Ref":b['business_ref'],"Product":b['product'],"Loading Point":b['loading_point'],"Trucks":f"{n}/{b['expected_trucks']}","Manifests":f"{m}/{b['expected_trucks']}","Booking Date":b['booking_date']})
        st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)
    else:
        st.info("Start by creating your first booking. Your saved bookings will appear here.")

elif page == "Create a booking":
    hero("📝 Create a booking", "Enter the shared details once. The booking date is set automatically on the first truck-list upload and will not change for later batches.")
    st.markdown('<div class="booking-form-marker"></div>', unsafe_allow_html=True)
    st.caption("The booking template and truck-list template are bundled with this prototype. Their columns are not edited by the form.")
    with st.form("create_booking_form"):
        c1,c2,c3=st.columns(3)
        with c1: booking_no=st.text_input("Booking Number *",placeholder="e.g. BEW600035000")
        with c2: business_ref=st.text_input("Business Reference *",placeholder="e.g. OPSE26100002")
        with c3: expected=st.number_input("Expected number of trucks *",min_value=1,max_value=1000,value=20,step=1)
        c4,c5,c6=st.columns(3)
        with c4: vessel=st.text_input("Vessel",placeholder="Vessel name")
        with c5: etb=st.text_input("ETB",placeholder="Date/time or leave blank")
        with c6: loading=st.text_input("Loading Point / Mine *",placeholder="e.g. MUTOKO")
        product=st.selectbox("ROAD MANIFEST category *",list(PRODUCTS.keys()),index=1)
        c7,c8=st.columns(2)
        with c7: description=st.text_input("Description of goods",value=product)
        with c8: weight=st.text_input("Weight / Mass",placeholder="e.g. 28 MT")
        c9,c10=st.columns(2)
        with c9: consignor=st.text_input("Consignor",placeholder="Consignor name")
        with c10: consignee=st.text_input("Consignee",placeholder="Consignee name")
        save=st.form_submit_button("✨ Save booking",type="primary",use_container_width=True)
    if save:
        no=clean(booking_no).upper()
        if not no or not clean(business_ref) or not clean(loading): st.error("Please complete Booking Number, Business Reference and Loading Point.")
        elif get_booking(no): st.error("That booking number already exists. Open it under Booking records instead of creating it again.")
        else:
            con.execute("INSERT INTO bookings VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",(no,clean(business_ref).upper(),clean(vessel),clean(etb),clean(loading).upper(),int(expected),"",product,clean(description) or product,clean(consignor),clean(consignee),clean(weight),now_local().isoformat(timespec="seconds")))
            con.commit()
            st.success(f"Booking {no} saved. Its booking date will be fixed automatically when you upload the first truck list.")
            st.info("Next: open ‘Add trucks & generate’ and upload the first truck batch.")

elif page == "Add trucks & generate":
    hero("🚛 Add trucks & generate documents", "Upload a new batch. The system keeps the first booking date, checks container formats, and creates the matching Excel truck list and Word manifests.")
    bookings=get_bookings()
    if not bookings:
        st.info("First create a booking so the uploaded trucks can be saved under the correct booking number.")
        if st.button("＋ Create your first booking", type="primary", use_container_width=True):
            st.session_state.nav_target = "Create a booking"
            st.rerun()
        st.markdown("You can still test whether your truck file uploads and its columns are readable under **Templates & setup → Test a truck-list upload**.")
    else:
        labels=[f"{b['booking_no']} · {b['loading_point']} · target {b['expected_trucks']}" for b in bookings]
        chosen=st.selectbox("Choose booking",range(len(bookings)),format_func=lambda i:labels[i])
        booking=bookings[chosen]
        existing=get_trucks(booking['booking_no'])
        c1,c2,c3=st.columns(3)
        with c1: metric("Expected trucks",booking['expected_trucks'],"Booking target")
        with c2: metric("Trucks received",len(existing),"Unique truck registrations")
        with c3: metric("Manifests generated",sum(int(t['manifest_generated']) for t in existing),"For this booking")
        uploaded=st.file_uploader("Upload truck details list",type=["xls","xlsx","xlsm","csv"],key="truck_batch_upload")
        if uploaded:
            try:
                df=parse_truck_file(uploaded.getvalue(),uploaded.name)
                st.markdown("#### Batch preview")
                st.dataframe(df.drop(columns=["_Transporter"],errors="ignore"),use_container_width=True,hide_index=True)
                if st.button("Validate and add this batch",type="primary",use_container_width=True):
                    added=[]; duplicates=[]; invalid=[]; missing=[]
                    current_max=max([int(t['sn']) for t in existing if t['sn'] is not None] or [0])
                    used_sns={int(t['sn']) for t in existing if t['sn'] is not None}
                    for _, row in df.iterrows():
                        proposed=source_value(row,"SN")
                        try: sn=int(float(proposed)) if proposed else current_max+1
                        except Exception: sn=current_max+1
                        while sn in used_sns and not proposed: sn+=1
                        rec=truck_record(row,booking['booking_no'],sn)
                        if not rec['truck_reg']: continue
                        matched = next((t for t in existing if t['truck_reg'].upper()==rec['truck_reg'].upper()), None)
                        if matched:
                            # A later booking list often adds container/seal data to trucks already saved
                            # from the first truck list. Match by normalized truck registration and merge
                            # non-empty fields rather than skipping the row as a duplicate.
                            incoming_container_status = validate_container(rec["container_no"])
                            if incoming_container_status == "invalid": invalid.append(f"{rec['truck_reg']}: container {rec['container_no']} must be 4 letters + 7 digits.")
                            if incoming_container_status == "missing": missing.append(rec["truck_reg"])
                            updates = {
                                "trailer_reg": rec["trailer_reg"], "container_no": rec["container_no"], "seal": rec["seal"],
                                "driver_name": rec["driver_name"], "phone_no": rec["phone_no"], "licence_no": rec["licence_no"],
                                "passport_no": rec["passport_no"], "capacity": rec["capacity"], "tare_weight": rec["tare_weight"],
                                "transporter": rec["transporter"]
                            }
                            updates = {k:v for k,v in updates.items() if clean(v)}
                            if updates:
                                set_sql = ", ".join(f"{k}=?" for k in updates)
                                con.execute(f"UPDATE trucks SET {set_sql} WHERE booking_no=? AND truck_reg=?", (*updates.values(), booking['booking_no'], rec['truck_reg']))
                                matched_row = dict(matched)
                                matched_row.update(updates)
                                added.append(matched_row)
                            else:
                                duplicates.append(rec['truck_reg'])
                            continue
                        if any(t['truck_reg'].upper()==rec['truck_reg'].upper() for t in added):
                            # If the same registration appears twice in this uploaded file, merge the later non-empty details.
                            prior = next(t for t in added if t['truck_reg'].upper()==rec['truck_reg'].upper())
                            for field in ("trailer_reg","container_no","seal","driver_name","phone_no","licence_no","passport_no","capacity","tare_weight","transporter"):
                                if clean(rec.get(field)): prior[field]=rec[field]
                            continue
                        if sn in used_sns:
                            invalid.append(f"{rec['truck_reg']}: SN {sn} is already used; reassigned to {current_max+1}.")
                            sn=current_max+1
                            while sn in used_sns: sn+=1
                            rec['sn']=sn
                        used_sns.add(sn); current_max=max(current_max,sn)
                        status=validate_container(rec['container_no'])
                        if status=="invalid": invalid.append(f"{rec['truck_reg']}: container {rec['container_no']} must be 4 letters + 7 digits.")
                        if status=="missing": missing.append(rec['truck_reg'])
                        added.append(rec)
                    # The first truck-list upload fixes the booking date; later batches never change it.
                    if added and not clean(booking['booking_date']):
                        fixed_date = now_local().strftime("%d.%m.%y")
                        con.execute("UPDATE bookings SET booking_date=? WHERE booking_no=?", (fixed_date, booking['booking_no']))
                        booking = get_booking(booking['booking_no'])
                    existing_regs = {t['truck_reg'].upper() for t in existing}
                    for rec in added:
                        if rec['truck_reg'].upper() in existing_regs:
                            continue  # already merged into the saved truck record above
                        con.execute("INSERT INTO trucks (booking_no,sn,truck_reg,trailer_reg,container_no,seal,driver_name,phone_no,licence_no,passport_no,capacity,tare_weight,transporter,batch_uploaded_at,manifest_generated) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,0)",
                          (rec['booking_no'],rec['sn'],rec['truck_reg'],rec['trailer_reg'],rec['container_no'],rec['seal'],rec['driver_name'],rec['phone_no'],rec['licence_no'],rec['passport_no'],rec['capacity'],rec['tare_weight'],rec['transporter'],now_local().isoformat(timespec="seconds")))
                    con.commit()
                    st.session_state['last_batch_booking']=booking['booking_no']
                    st.session_state['last_batch_sns']=sorted(set(int(r['sn']) for r in added))
                    st.session_state['last_batch_trucks']=sorted(set(r['truck_reg'] for r in added))
                    st.session_state['last_batch_warning']={"duplicates":duplicates,"invalid":invalid,"missing":missing}
                    if added: st.success(f"Processed {len(added)} truck row(s) for {booking['booking_no']}. New trucks were added and matching saved trucks were updated with any supplied details.")
                    else: st.warning("No new trucks were added. Check duplicates and the uploaded list.")
                    if duplicates: st.warning("Duplicate registrations skipped: " + ", ".join(duplicates))
                    if invalid: st.warning("Review these records:\n\n"+"\n".join(invalid))
                    if missing: st.info("Container number missing for: "+", ".join(missing))
                    st.rerun()
            except Exception as e: st.error(f"Could not read this truck list: {e}")
        # Downloads for latest batch in session; also provide full booking list.
        if st.session_state.get('last_batch_booking')==booking['booking_no'] and st.session_state.get('last_batch_sns'):
            latest_sns=st.session_state['last_batch_sns']
            latest_all=get_trucks(booking['booking_no'])
            latest=[t for t in latest_all if int(t['sn']) in latest_sns]
            latest.sort(key=lambda t:int(t['sn']))
            lo,hi=min(latest_sns),max(latest_sns)
            st.markdown("#### Your generated files")
            a,b,c=st.columns(3)
            with a:
                tl=make_truck_list(booking,latest)
                fname=f"{booking['booking_no']} TRUCK LIST ({lo}-{hi}).xlsx"
                st.download_button("📊 Download batch truck list",tl,file_name=fname,mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
            with b:
                all_bytes=make_booking_list(booking,latest_all)
                st.download_button("📋 Download full booking list",all_bytes,file_name=f"{booking['booking_no']} BOOKING LIST.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
            with c:
                if st.button("📄 Generate Road Manifests",type="primary",use_container_width=True):
                    zbuf,count,issues=make_manifest_zip(booking,latest)
                    st.session_state['last_manifest_zip']=zbuf
                    st.session_state['last_manifest_name']=f"{booking['booking_no']} ROAD MANIFESTS ({lo}-{hi}).zip"
                    st.session_state['last_manifest_booking']=booking['booking_no']
                    st.session_state['last_manifest_sns']=latest_sns
                    st.session_state['last_manifest_issues']=issues
                    st.session_state['last_manifest_count']=count
                    con.executemany("UPDATE trucks SET manifest_generated=1 WHERE booking_no=? AND sn=?",[(booking['booking_no'],sn) for sn in latest_sns])
                    con.commit(); st.rerun()
        if st.session_state.get('last_manifest_booking')==booking['booking_no'] and st.session_state.get('last_manifest_zip'):
            st.success(f"Prepared {st.session_state.get('last_manifest_count',0)} Word manifest(s).")
            st.download_button("⬇️ Download Road Manifests (ZIP)",st.session_state['last_manifest_zip'],file_name=st.session_state['last_manifest_name'],mime="application/zip",use_container_width=True)
            issues=st.session_state.get('last_manifest_issues',[])
            if issues: st.warning("Please review before official use:\n\n"+"\n".join(issues))
        if existing:
            st.markdown("#### Trucks currently recorded for this booking")
            view=[]
            for t in get_trucks(booking['booking_no']):
                view.append({"SN":t['sn'],"Truck Reg.":t['truck_reg'],"Trailer Reg.":t['trailer_reg'],"Container No.":t['container_no'],"Transporter":t['transporter'],"Driver":t['driver_name'],"Manifest":"Generated" if t['manifest_generated'] else "Not generated"})
            st.dataframe(pd.DataFrame(view),use_container_width=True,hide_index=True)

elif page == "Booking records":
    hero("📚 Booking records", "Open a booking to check progress. Your counts are based on saved truck records, not just downloaded files.")
    bookings=get_bookings()
    if not bookings: st.info("No bookings yet. Create one to begin.")
    for b in bookings:
        trucks=get_trucks(b['booking_no']); n=len(trucks); m=sum(int(t['manifest_generated']) for t in trucks)
        with st.expander(f"{b['booking_no']} · {b['product']} · {n}/{b['expected_trucks']} trucks · {m}/{b['expected_trucks']} manifests"):
            c1,c2,c3=st.columns(3)
            with c1: st.write("**Business Ref:**",b['business_ref']); st.write("**Vessel:**",b['vessel'] or "—")
            with c2: st.write("**Loading Point:**",b['loading_point']); st.write("**Booking Date:**",b['booking_date'])
            with c3: st.write("**Expected trucks:**",b['expected_trucks']); st.write("**ETB:**",b['etb'] or "—")
            st.write(f"**Consignor:** {b['consignor'] or '—'} · **Consignee:** {b['consignee'] or '—'} · **Weight/Mass:** {b['weight_mass'] or '—'}")
            st.progress(min(n/max(int(b['expected_trucks']),1),1.0),text=f"Trucks received: {n}/{b['expected_trucks']}")
            st.progress(min(m/max(int(b['expected_trucks']),1),1.0),text=f"Manifests generated: {m}/{b['expected_trucks']}")
            if trucks:
                st.dataframe(pd.DataFrame([{"SN":t['sn'],"Truck":t['truck_reg'],"Container":t['container_no'],"Manifest":"Generated" if t['manifest_generated'] else "Pending"} for t in trucks]),use_container_width=True,hide_index=True)

elif page == "Templates & setup":
    hero("🎨 Templates & setup", "A colourful, responsive workspace with reusable templates and clear validation rules.")
    st.markdown("### Test a truck-list upload")
    st.write("Test the uploader here before creating a booking. This preview does not save truck records or generate official documents.")
    test_upload = st.file_uploader("Choose a truck details file to test", type=["xls", "xlsx", "xlsm", "csv"], key="setup_upload_test")
    if test_upload is not None:
        try:
            test_df = parse_truck_file(test_upload.getvalue(), test_upload.name)
            st.success(f"Upload successful — found {len(test_df)} truck row(s).")
            st.dataframe(test_df.drop(columns=["_Transporter"], errors="ignore"), use_container_width=True, hide_index=True)
            st.caption("If the preview looks correct, create a booking and upload this file under Add trucks & generate to save the trucks.")
        except Exception as exc:
            st.error(f"The file uploaded, but its columns could not be read: {exc}")
            st.info("Supported formats: .xls, .xlsx, .xlsm and .csv. The truck list must contain Truck Reg. and Trailer/Trailor Reg. column headings.")
    st.markdown("### Bundled templates")
    st.success("Booking List template saved in assets/templates/booking_template.xlsx")
    st.success("Truck List template saved in assets/templates/truck_list_template.xlsx")
    st.success("ROAD MANIFEST base template saved, with four category-specific copies.")
    st.warning("The four category files currently begin as copies of the single ROAD MANIFEST template you supplied. If any product has a different official layout, upload that specific template before production use.")
    st.markdown("### Rules included in this prototype")
    st.markdown("- Truck and trailer registrations are normalized without spaces.\n- OP FLEET A → OP (A); OP FLEET B → OP (B).\n- Booking date is set once on the first truck-list upload, using Africa/Harare local date.\n- Container number must match **4 letters + 7 digits**; missing and invalid values are flagged.\n- Batch truck-list filenames use booking number and sequence range, e.g. `BEW600035000 TRUCK LIST (17-20).xlsx`.\n- The original Excel column order is retained.")
    st.markdown("### Online privacy and persistence")
    st.error("This is a development prototype. Its SQLite database is stored on the app's local disk, which is not reliable permanent storage on many cloud hosts. Before using it online with real driver or customs data, connect a persistent private database, configure access control, and test backups.")
    st.code('APP_PASSWORD = "use-a-long-unique-password"\n# Add this to Streamlit app Secrets for a basic password gate.',language="toml")
    st.caption("Do not put real truck lists, manifests, passwords, or driver details in a public GitHub repository.")
